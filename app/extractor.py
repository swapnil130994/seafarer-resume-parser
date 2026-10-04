import os
import io
import fitz
import pdfplumber
from docx import Document
from PIL import Image
import pytesseract


def normalize_cell_text(value):
    value = str(value or "").replace("\xa0", " ")
    return " ".join(value.split())


def extract_docx(path):
    doc = Document(path)
    parts = []

    for paragraph in doc.paragraphs:
        # Preserve explicit line breaks inside Word paragraphs. Some CVs put
        # the role/name/contact fields into one paragraph with visual breaks.
        raw_lines = paragraph.text.splitlines() or [paragraph.text]
        for raw_line in raw_lines:
            value = normalize_cell_text(raw_line)
            if value:
                parts.append(value)

    for table in doc.tables:
        for row in table.rows:
            cells = [normalize_cell_text(cell.text) for cell in row.cells]
            if any(cells):
                # Keep table columns separated. The parser uses this structure
                # for employment, documents and course tables.
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def ocr_image(image):
    return pytesseract.image_to_string(image, lang="eng")


def extract_pdf(path):
    parts = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                if text.strip():
                    parts.append(text)

                # Preserve PDF tables where pdfplumber can detect them.
                try:
                    tables = page.extract_tables() or []
                    for table in tables:
                        for row in table:
                            cells = [normalize_cell_text(cell) for cell in (row or [])]
                            if any(cells):
                                parts.append(" | ".join(cells))
                except Exception:
                    pass
    except Exception:
        pass

    text = "\n".join(parts).strip()
    if len(text) < 150:
        output = []
        doc = fitz.open(path)
        try:
            for page in doc:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                output.append(ocr_image(image))
        finally:
            doc.close()
        return "\n".join(output).strip(), True

    return text, False


def extract_text(path):
    extension = os.path.splitext(path)[1].lower()
    if extension == ".pdf":
        return extract_pdf(path)
    if extension == ".docx":
        return extract_docx(path), False
    if extension in {".jpg", ".jpeg", ".png", ".webp"}:
        return ocr_image(Image.open(path)), True
    if extension == ".doc":
        raise ValueError("Old .doc files are not supported directly. Convert .doc to .docx or PDF.")
    raise ValueError("Unsupported file type.")
